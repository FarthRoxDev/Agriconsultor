#!/usr/bin/env python3
"""presentacion-python-pptx.py — genera un .pptx data-driven con python-pptx.

Uso:   python3 entregables/plantillas/presentacion-python-pptx.py
Salida: entregables/salidas/presentacion-python-pptx.pptx
Regla CLAUDE.md: toda cifra conserva su fuente (ver FUENTE).
"""
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE

FUENTE = "MAG 2016, cit. en INTA 2016 (Manual del cultivo de papa en CR), p. 11"
rendimiento = [
    ("Cartago", 25.5), ("Alvarado", 25.2), ("Zarcero", 23.8), ("Turrialba", 23.8),
    ("Naranjo", 23.6), ("El Guarco", 23.1), ("Oreamuno", 22.3), ("Dota", 21.5),
]

prs = Presentation()

# Portada
s1 = prs.slides.add_slide(prs.slide_layouts[0])
s1.shapes.title.text = "Papa en Costa Rica"
s1.placeholders[1].text = "Rendimiento por cantón — Agriconsultor"

# Gráfico data-driven
s2 = prs.slides.add_slide(prs.slide_layouts[5])
s2.shapes.title.text = "Rendimiento por cantón (t/ha)"
data = CategoryChartData()
data.categories = [c for c, _ in rendimiento]
data.add_series("t/ha", [v for _, v in rendimiento])
s2.shapes.add_chart(
    XL_CHART_TYPE.COLUMN_CLUSTERED,
    Inches(0.5), Inches(1.5), Inches(9), Inches(4.5), data,
)
tb = s2.shapes.add_textbox(Inches(0.5), Inches(6.2), Inches(9), Inches(0.4))
p = tb.text_frame.paragraphs[0]
p.text = f"Fuente: {FUENTE}"
p.font.size = Pt(10)
p.font.italic = True

out = Path(__file__).resolve().parents[1] / "salidas"
out.mkdir(parents=True, exist_ok=True)
dest = out / "presentacion-python-pptx.pptx"
prs.save(str(dest))
print("OK ->", dest)
